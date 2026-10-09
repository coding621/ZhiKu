#!/usr/bin/env python3
"""知库 ZhiKu —— 本地知识库桌面工具。

运行模式：
    python kb-server.py            浏览器模式（开发调试）
    python kb-server.py --gui      独立窗口模式（pywebview + WebView2）
    打包成 exe 后双击默认即为独立窗口模式（--no-window 可退回浏览器）

数据与配置存放在：
    开发：脚本所在目录
    exe ：%LOCALAPPDATA%\\ZhiKu\\
"""
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import os
import re
import sqlite3
import sys
import threading
import time
import traceback
import urllib.parse
from html.parser import HTMLParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

APP_NAME = "知库 ZhiKu"
APP_REG = "ZhiKuKB"          # 注册表自启动项名
MAX_ROOTS = 100

TEXT_EXTS = {".md", ".markdown", ".txt", ".log", ".csv", ".json", ".yml", ".yaml", ".xml", ".rst", ".tex"}
HTML_EXTS = {".html", ".htm"}
DOCX_EXTS = {".docx"}
PDF_EXTS = {".pdf"}
XLSX_EXTS = {".xlsx", ".xlsm"}
PPTX_EXTS = {".pptx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp", ".ico"}
INDEXABLE = TEXT_EXTS | HTML_EXTS | DOCX_EXTS | PDF_EXTS | XLSX_EXTS | PPTX_EXTS
ALL_EXTS = INDEXABLE | IMAGE_EXTS | {".doc", ".xls", ".ppt", ".zip", ".7z", ".rar", ".gz", ".mp4", ".mp3"}

_is_frozen = bool(getattr(sys, "frozen", False))


# ---------------------------------------------------------------- 路径与配置

def app_dir() -> Path:
    """exe 模式放用户数据目录，开发模式放脚本目录。"""
    if _is_frozen:
        d = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "ZhiKu"
        d.mkdir(parents=True, exist_ok=True)
        return d
    return Path(__file__).resolve().parent


DATA_DIR = app_dir()
CONFIG_PATH = DATA_DIR / "config.json"
DB_PATH = DATA_DIR / "index.db"
DEFAULT_ROOT = DATA_DIR / "kb"
if _is_frozen:
    BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", str(DATA_DIR)))
else:
    BUNDLE_DIR = Path(__file__).resolve().parent
STATIC_FILE = BUNDLE_DIR / "static" / "index.html"

_cfg_lock = threading.RLock()
_config: dict | None = None


def _default_config() -> dict:
    return {"port": 8360, "autostart": False, "roots": [str(DEFAULT_ROOT)]}


def config() -> dict:
    global _config
    with _cfg_lock:
        if _config is None:
            cfg = _default_config()
            if CONFIG_PATH.exists():
                try:
                    loaded = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
                    if isinstance(loaded, dict):
                        cfg.update(loaded)
                except (json.JSONDecodeError, OSError):
                    pass
            cfg.setdefault("port", 8360)
            cfg.setdefault("autostart", False)
            roots = cfg.get("roots")
            if not isinstance(roots, list):
                roots = []
            cfg["roots"] = [str(r) for r in roots][:MAX_ROOTS]
            _config = cfg
        return _config


def save_config() -> None:
    with _cfg_lock:
        CONFIG_PATH.write_text(
            json.dumps(_config, ensure_ascii=False, indent=2), encoding="utf-8"
        )


def norm_root(p: str) -> str:
    return str(Path(str(p)).expanduser().resolve())


def root_id(p: str) -> str:
    key = os.path.normcase(norm_root(p))
    return hashlib.md5(key.encode("utf-8")).hexdigest()[:12]


def roots_list() -> list[str]:
    return [norm_root(r) for r in config()["roots"]]


def roots_map() -> dict[str, str]:
    """{root_id: 绝对路径}"""
    return {root_id(r): norm_root(r) for r in config()["roots"]}


def root_names() -> dict[str, str]:
    """{root_id: 显示名}（重名时加短后缀区分）"""
    result: dict[str, str] = {}
    seen: dict[str, int] = {}
    for r in roots_list():
        name = Path(r).name or r
        seen[name] = seen.get(name, 0) + 1
    used: dict[str, int] = {}
    for r in roots_list():
        rid = root_id(r)
        base = Path(r).name or r
        if seen[base] > 1:
            used[base] = used.get(base, 0) + 1
            result[rid] = f"{base}·{rid[:4]}"
        else:
            result[rid] = base
    return result


def seed_default_kb() -> None:
    """首次运行（尤其 exe）给默认知识库放一篇使用说明。"""
    if DEFAULT_ROOT.exists() and any(DEFAULT_ROOT.iterdir()):
        return
    DEFAULT_ROOT.mkdir(parents=True, exist_ok=True)
    (DEFAULT_ROOT / "快速上手.md").write_text(
        "# 快速上手\n\n"
        "这是你的本地知识库。\n\n"
        "## 添加内容\n\n"
        "1. 在「设置」里添加任意电脑上的文件夹作为知识库\n"
        "2. 把文档丢进文件夹（支持 md / txt / html / pdf / docx / xlsx / pptx）\n"
        "3. 点侧栏「重新扫描」即可被搜索到\n\n"
        "## 快捷键\n\n"
        "- `/` 聚焦搜索框\n"
        "- `Esc` 关闭弹窗\n",
        encoding="utf-8",
    )


# ---------------------------------------------------------------- 文本抽取

class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in ("p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"):
            self.parts.append("\n")

    def handle_data(self, data):
        if self._skip:
            return
        if self._in_title:
            self.title += data
        else:
            self.parts.append(data)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    # 1) BOM 优先（Windows 记事本常见的 UTF-16/UTF-32/UTF-8 带签名）
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return raw.decode("utf-32")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    # 2) 无 BOM 的 UTF-16 启发式：样本里 \x00 占比异常高
    sample = raw[:4096]
    if len(raw) >= 2 and len(raw) % 2 == 0 and len(sample) >= 16:
        zeros = sample.count(0)
        if zeros > len(sample) // 4:
            odd = sum(1 for i in range(1, len(sample), 2) if sample[i] == 0)
            even = sum(1 for i in range(0, len(sample), 2) if sample[i] == 0)
            try:
                return raw.decode("utf-16-le" if odd >= even else "utf-16-be")
            except UnicodeDecodeError:
                pass
    # 3) 常规编码
    for enc in ("utf-8", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _extract(path: Path, ext: str) -> str:
    if ext in TEXT_EXTS:
        return _read_text(path)
    if ext in HTML_EXTS:
        parser = _HTMLText()
        parser.feed(_read_text(path))
        title = parser.title.strip()
        body = re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()
        return f"# {title}\n\n{body}" if title else body
    if ext in DOCX_EXTS:
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext in PDF_EXTS:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext in XLSX_EXTS:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)
    if ext in PPTX_EXTS:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"## 幻灯片 {i}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs).strip()
                        if text:
                            parts.append(text)
        return "\n".join(parts)
    return ""


def _make_title(path: Path, ext: str, text: str) -> str:
    if ext in (".md", ".markdown", ".html", ".htm"):
        for line in text.splitlines()[:10]:
            m = re.match(r"^#\s+(.+)$", line.strip())
            if m:
                return m.group(1).strip()
    return path.stem


# ---------------------------------------------------------------- 索引

_lock = threading.RLock()
_conn: sqlite3.Connection | None = None
INDEX_VERSION = 2          # 抽取/解码逻辑变更时 +1，强制重建索引


def db() -> sqlite3.Connection:
    global _conn
    if _conn is None:
        _conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        _conn.row_factory = sqlite3.Row
        _conn.execute("PRAGMA journal_mode=WAL")
        _conn.execute("CREATE TABLE IF NOT EXISTS kv (k TEXT PRIMARY KEY, v TEXT)")
        ver = _conn.execute("SELECT v FROM kv WHERE k='index_version'").fetchone()
        if ver is None or ver["v"] != str(INDEX_VERSION):
            _conn.execute("DROP TABLE IF EXISTS docs")
            _conn.execute(
                "INSERT OR REPLACE INTO kv(k,v) VALUES ('index_version', ?)",
                (str(INDEX_VERSION),),
            )
            _conn.commit()
        cols = {r["name"] for r in _conn.execute("PRAGMA table_info(docs)")}
        if cols and "root_id" not in cols:
            _conn.execute("DROP TABLE docs")            # 旧版单库结构 → 重建
            cols = set()
        _conn.execute("""
            CREATE TABLE IF NOT EXISTS docs (
                pk TEXT PRIMARY KEY,
                root_id TEXT NOT NULL,
                path TEXT NOT NULL,
                sub TEXT NOT NULL,
                title TEXT NOT NULL,
                ext TEXT NOT NULL,
                mtime REAL NOT NULL,
                size INTEGER NOT NULL,
                content TEXT NOT NULL
            )
        """)
        _conn.execute("CREATE INDEX IF NOT EXISTS idx_docs_root ON docs(root_id)")
        _conn.execute("CREATE INDEX IF NOT EXISTS idx_docs_mtime ON docs(mtime DESC)")
        _conn.commit()
    return _conn


def _sub_of(rel: str) -> str:
    return rel.split("/", 1)[0] if "/" in rel else ""


def index_file(rid: str, rel: str, abs_path: Path) -> bool:
    ext = abs_path.suffix.lower()
    if ext not in INDEXABLE:
        return False
    try:
        text = _extract(abs_path, ext)
        st = abs_path.stat()
        with _lock:
            db().execute(
                "INSERT OR REPLACE INTO docs(pk,root_id,path,sub,title,ext,mtime,size,content) "
                "VALUES (?,?,?,?,?,?,?,?,?)",
                (f"{rid}/{rel}", rid, rel, _sub_of(rel),
                 _make_title(abs_path, ext, text), ext, st.st_mtime, st.st_size, text),
            )
            db().commit()
        return True
    except Exception:
        traceback.print_exc()
        return False


def rescan() -> dict:
    t0 = time.time()
    seen: set[str] = set()
    added = updated = skipped = failed = 0
    active_ids = set()
    for r in roots_list():
        rid = root_id(r)
        active_ids.add(rid)
        rootp = Path(r)
        if not rootp.is_dir():
            continue
        for dirpath, dirnames, filenames in os.walk(rootp):
            dirnames[:] = [d for d in dirnames if not d.startswith(".") and not d.startswith("~")]
            for name in sorted(filenames):
                if name.startswith(".") or name.startswith("~$"):
                    continue
                ap = Path(dirpath) / name
                ext = ap.suffix.lower()
                if ext not in ALL_EXTS:
                    continue
                rel = ap.relative_to(rootp).as_posix()
                pk = f"{rid}/{rel}"
                seen.add(pk)
                if ext not in INDEXABLE:
                    continue
                try:
                    st = ap.stat()
                except OSError:
                    continue
                with _lock:
                    row = db().execute(
                        "SELECT mtime,size FROM docs WHERE pk=?", (pk,)
                    ).fetchone()
                if row and abs(row["mtime"] - st.st_mtime) < 1e-6 and row["size"] == st.st_size:
                    skipped += 1
                    continue
                if index_file(rid, rel, ap):
                    updated += 1 if row else 0
                    added += 0 if row else 1
                else:
                    failed += 1
    with _lock:
        if seen:
            marks = ",".join("?" * len(seen))
            db().execute(f"DELETE FROM docs WHERE pk NOT IN ({marks})", list(seen))
        else:
            db().execute("DELETE FROM docs")
        if active_ids:
            ids = ",".join("?" * len(active_ids))
            db().execute(f"DELETE FROM docs WHERE root_id NOT IN ({ids})", list(active_ids))
        db().commit()
    return {"added": added, "updated": updated, "unchanged": skipped,
            "failed": failed, "took_ms": int((time.time() - t0) * 1000)}

# ---------------------------------------------------------------- 搜索

def _escape_like(term: str) -> str:
    return term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _snippet(content: str, terms: list[str], width: int = 170) -> str:
    low = content.lower()
    positions = [p for p in (low.find(t.lower()) for t in terms) if p >= 0]
    pos = min(positions) if positions else 0
    start = max(0, pos - width // 3)
    window = content[start:start + width + (len(terms[0]) if terms else 0)]
    window = re.sub(r"\s+", " ", window).strip()
    window = re.sub(r"\s*#{1,6}\s*", " ", window)
    window = window.replace("`", "").replace("**", "").replace("__", "")
    window = re.sub(r"\s{2,}", " ", window).strip()
    prefix = "…" if start > 0 else ""
    suffix = "…" if start + width < len(content) else ""
    return prefix + window + suffix


def _display(pk: str, names: dict[str, str]) -> str:
    rid, _, rel = pk.partition("/")
    return f"{names.get(rid, rid)}/{rel}" if rel else names.get(rid, rid)


def _group(row: sqlite3.Row, names: dict[str, str]) -> str:
    base = names.get(row["root_id"], row["root_id"])
    return f"{base}/{row['sub']}" if row["sub"] else base


def _row_dict(row: sqlite3.Row, snippet: str, names: dict[str, str]) -> dict:
    return {
        "path": row["pk"],
        "title": row["title"],
        "category": _group(row, names),
        "group": _group(row, names),
        "display": _display(row["pk"], names),
        "sub": row["sub"],
        "ext": row["ext"],
        "mtime": row["mtime"],
        "size": row["size"],
        "snippet": snippet,
    }


def search(q: str, root: str = "", sub: str = "", limit: int = 60) -> list[dict]:
    q = (q or "").strip()
    names = root_names()
    scope, scope_params = [], []
    if root:
        scope.append("root_id=?")
        scope_params.append(root)
        if sub:
            scope.append("sub=?")
            scope_params.append(sub)
    with _lock:
        if not q:
            sql = "SELECT * FROM docs"
            if scope:
                sql += " WHERE " + " AND ".join(scope)
            sql += " ORDER BY mtime DESC LIMIT ?"
            rows = db().execute(sql, scope_params + [limit]).fetchall()
            return [_row_dict(r, _snippet(r["content"], []), names) for r in rows]

        terms = [t for t in re.split(r"\s+", q) if t]
        conds = list(scope)
        params = list(scope_params)
        for t in terms:
            conds.append("(content LIKE ? ESCAPE '\\' OR title LIKE ? ESCAPE '\\')")
            pat = f"%{_escape_like(t)}%"
            params += [pat, pat]
        rows = db().execute(
            f"SELECT * FROM docs WHERE {' AND '.join(conds)}", params
        ).fetchall()

    scored: list[tuple[float, sqlite3.Row]] = []
    for row in rows:
        lc = row["content"].lower()
        title_l = row["title"].lower()
        score = 0.0
        for t in terms:
            lt = t.lower()
            hits = lc.count(lt)
            if hits == 0 and lt not in title_l:
                score = -1
                break
            score += min(hits, 20)
            if lt in title_l:
                score += 25
        if score >= 0:
            score += min(row["mtime"] / 1e9, 10) * 0.1
            scored.append((score, row))
    scored.sort(key=lambda x: (-x[0], -x[1]["mtime"]))
    return [_row_dict(r, _snippet(r["content"], terms), names) for _, r in scored[:limit]]


def get_doc(pk: str) -> dict | None:
    with _lock:
        row = db().execute("SELECT * FROM docs WHERE pk=?", (pk,)).fetchone()
    if row is None:
        return None
    names = root_names()
    d = _row_dict(row, "", names)
    d["content"] = row["content"]
    d["raw_url"] = "/api/raw?path=" + urllib.parse.quote(pk)
    return d


def safe_path(pk: str) -> Path | None:
    """pk = root_id/相对路径 → 绝对路径；阻止越出所属根目录。"""
    if not pk or "/" not in pk:
        return None
    rid, _, rel = pk.partition("/")
    root = roots_map().get(rid)
    if not root:
        return None
    try:
        target = (Path(root) / rel).resolve()
        rootp = Path(root).resolve()
        if rootp not in target.parents:
            return None
        return target if target.is_file() else None
    except (OSError, ValueError):
        return None


# ---------------------------------------------------------------- 元信息 / 笔记

def meta() -> dict:
    names = root_names()
    roots_info, folders = [], []
    total = 0
    with _lock:
        for r in roots_list():
            rid = root_id(r)
            rows = db().execute(
                "SELECT sub, COUNT(*) AS n FROM docs WHERE root_id=? GROUP BY sub", (rid,)
            ).fetchall()
            count = sum(r_["n"] for r_ in rows)
            total += count
            subs = sorted(
                ({"name": r_["sub"], "count": r_["n"]} for r_ in rows if r_["sub"]),
                key=lambda x: -x["count"],
            )
            exists = Path(r).is_dir()
            roots_info.append({
                "id": rid, "path": r, "name": names[rid],
                "exists": exists, "count": count, "subs": subs,
            })
            folders.append({"id": rid, "label": names[rid], "root_id": rid, "sub": ""})
            for s in subs:
                folders.append({
                    "id": f"{rid}|{s['name']}",
                    "label": f"{names[rid]} / {s['name']}",
                    "root_id": rid, "sub": s["name"],
                })
    return {
        "roots": roots_info,
        "folders": folders,
        "total": total,
        "autostart": get_autostart(),
        "port": config().get("port", 8360),
        "data_dir": str(DATA_DIR),
        "max_roots": MAX_ROOTS,
        "frozen": _is_frozen,
    }


def create_note(folder_id: str, title: str, body: str) -> dict:
    title = (title or "").strip() or "未命名笔记"
    rid, _, sub = (folder_id or "").partition("|")
    root = roots_map().get(rid)
    if not root:
        raise ValueError("目标知识库不存在，请先在设置中添加")
    folder = Path(root)
    if sub:
        folder = folder / sub
        if not folder.is_dir():
            raise ValueError(f"子目录不存在：{sub}")
    safe = re.sub(r'[\\/:*?"<>|]', "_", title).strip(" .") or "未命名笔记"
    path = folder / f"{safe}.md"
    if path.exists():
        path = folder / f"{safe}-{time.strftime('%Y%m%d-%H%M%S')}.md"
    path.write_text(f"# {title}\n\n{(body or '').rstrip()}\n", encoding="utf-8")
    rel = path.relative_to(Path(root)).as_posix()
    index_file(rid, rel, path)
    names = root_names()
    return {"path": f"{rid}/{rel}", "display": f"{names.get(rid, rid)}/{rel}"}


# ---------------------------------------------------------------- 自启动（Windows）

RUN_KEY = r"Software\Microsoft\Windows\CurrentVersion\Run"


def _autostart_cmd() -> str:
    if _is_frozen:
        return f'"{sys.executable}"'
    return f'"{sys.executable}" "{Path(__file__).resolve()}" --gui'


def get_autostart() -> bool:
    if os.name != "nt":
        return False
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, RUN_KEY, 0, winreg.KEY_READ) as key:
            winreg.QueryValueEx(key, APP_REG)
            return True
    except OSError:
        return False


def set_autostart(on: bool) -> None:
    if os.name != "nt":
        raise RuntimeError("仅支持 Windows")
    import winreg
    with winreg.CreateKeyEx(
        winreg.HKEY_CURRENT_USER, RUN_KEY, 0, winreg.KEY_SET_VALUE
    ) as key:
        if on:
            winreg.SetValueEx(key, APP_REG, 0, winreg.REG_SZ, _autostart_cmd())
        else:
            try:
                winreg.DeleteValue(key, APP_REG)
            except FileNotFoundError:
                pass
    config()["autostart"] = on
    save_config()


def apply_settings(roots: list[str], autostart: bool) -> dict:
    """保存设置：目录列表（≤100）+ 自启动开关。"""
    warnings = []
    cleaned: list[str] = []
    seen: set[str] = set()
    if len(roots) > MAX_ROOTS:
        raise ValueError(f"最多支持 {MAX_ROOTS} 个知识库目录")
    for raw in roots:
        raw = str(raw or "").strip().strip('"')
        if not raw:
            continue
        try:
            p = norm_root(raw)
        except (OSError, ValueError):
            raise ValueError(f"无效路径：{raw}")
        key = os.path.normcase(p)
        if key in seen:
            continue
        if not Path(p).is_dir():
            raise ValueError(f"目录不存在：{p}")
        seen.add(key)
        cleaned.append(p)
    cfg = config()
    cfg["roots"] = cleaned[:MAX_ROOTS]
    save_config()
    try:
        if bool(autostart) != get_autostart():
            set_autostart(bool(autostart))
    except Exception as e:
        warnings.append(f"自启动设置失败：{e}")
    stats = rescan()
    result = meta()
    result["scan"] = stats
    result["warnings"] = warnings
    return result


# ---------------------------------------------------------------- HTTP

class Handler(BaseHTTPRequestHandler):
    server_version = "ZhiKu/2.0"

    def log_message(self, fmt, *args):
        pass

    def _json(self, obj, status: int = 200):
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def _bytes(self, data: bytes, ctype: str, status: int = 200, extra: dict | None = None):
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(data)

    def _file(self, path: Path, download: bool = False):
        ctype = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        if path.suffix.lower() in (".md", ".markdown"):
            ctype = "text/markdown; charset=utf-8"
        if ctype.startswith("text/") and "charset" not in ctype:
            ctype += "; charset=utf-8"
        data = path.read_bytes()
        fname = urllib.parse.quote(path.name)
        disp = "attachment" if download else "inline"
        self._bytes(data, ctype, extra={
            "Content-Disposition": f"{disp}; filename*=UTF-8''{fname}",
            "Cache-Control": "no-cache",
        })

    def _body_json(self) -> dict:
        length = int(self.headers.get("Content-Length") or 0)
        if length <= 0 or length > 20 * 1024 * 1024:
            return {}
        try:
            return json.loads(self.rfile.read(length).decode("utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError):
            return {}

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        qs = urllib.parse.parse_qs(parsed.query)
        route = parsed.path
        try:
            if route in ("/", "/index.html"):
                if not STATIC_FILE.exists():
                    self._bytes(b"static/index.html missing", "text/plain; charset=utf-8", 404)
                    return
                self._file(STATIC_FILE)
                return
            if route == "/api/meta":
                self._json(meta())
                return
            if route == "/api/settings":
                s = meta()
                s["roots_config"] = list(config()["roots"])
                self._json(s)
                return
            if route == "/api/search":
                self._json({"results": search(
                    qs.get("q", [""])[0],
                    qs.get("root", [""])[0],
                    qs.get("sub", [""])[0],
                )})
                return
            if route == "/api/doc":
                doc = get_doc(qs.get("path", [""])[0])
                if doc is None:
                    self._json({"error": "文档不存在"}, 404)
                    return
                self._json(doc)
                return
            if route == "/api/raw":
                target = safe_path(qs.get("path", [""])[0])
                if target is None:
                    self._json({"error": "非法路径"}, 403)
                    return
                self._file(target, download=qs.get("dl", ["0"])[0] == "1")
                return
            if route == "/api/rescan":
                result = rescan()
                result["meta"] = meta()
                self._json(result)
                return
            self._json({"error": "not found"}, 404)
        except BrokenPipeError:
            pass
        except Exception as e:
            traceback.print_exc()
            self._json({"error": str(e)}, 500)

    def do_POST(self):
        route = urllib.parse.urlparse(self.path).path
        try:
            data = self._body_json()
            if route == "/api/note":
                try:
                    note = create_note(
                        data.get("folder", ""), data.get("title", ""), data.get("body", "")
                    )
                except ValueError as e:
                    self._json({"error": str(e)}, 400)
                    return
                self._json(note)
                return
            if route == "/api/settings":
                roots = data.get("roots")
                if not isinstance(roots, list):
                    self._json({"error": "参数错误"}, 400)
                    return
                try:
                    result = apply_settings(roots, bool(data.get("autostart")))
                except ValueError as e:
                    self._json({"error": str(e)}, 400)
                    return
                self._json(result)
                return
            if route == "/api/rescan":
                self._json(rescan())
                return
            self._json({"error": "not found"}, 404)
        except BrokenPipeError:
            pass
        except Exception as e:
            traceback.print_exc()
            self._json({"error": str(e)}, 500)


# ---------------------------------------------------------------- 启动

def pick_server(preferred: int) -> ThreadingHTTPServer:
    last_err: OSError | None = None
    for port in range(preferred, preferred + 12):
        try:
            return ThreadingHTTPServer(("127.0.0.1", port), Handler)
        except OSError as e:
            last_err = e
    raise last_err  # type: ignore[misc]


def serve_in_thread(server: ThreadingHTTPServer) -> threading.Thread:
    t = threading.Thread(target=server.serve_forever, daemon=True)
    t.start()
    return t


def run_gui(url: str) -> None:
    import webview
    webview.create_window(
        f"{APP_NAME} · 本地知识库", url,
        width=1380, height=900, min_size=(980, 640),
        background_color="#F7F4EC",
    )
    webview.start()


def main() -> None:
    ap = argparse.ArgumentParser(description=APP_NAME)
    ap.add_argument("--port", type=int, default=None)
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--gui", action="store_true", help="独立窗口模式")
    ap.add_argument("--no-window", action="store_true", help="exe 下退回浏览器模式")
    ap.add_argument("--rescan-only", action="store_true")
    args = ap.parse_args()

    cfg = config()
    if args.port:
        cfg["port"] = args.port
        save_config()
    if not cfg["roots"]:
        cfg["roots"] = [str(DEFAULT_ROOT)]
        save_config()
    for r in list(cfg["roots"]):
        Path(r).mkdir(parents=True, exist_ok=True)
    seed_default_kb()

    stats = rescan()
    print(f"[{APP_NAME}] 索引：新增 {stats['added']} / 更新 {stats['updated']} / "
          f"跳过 {stats['unchanged']} / 失败 {stats['failed']}（{stats['took_ms']} ms）")
    if args.rescan_only:
        return

    server = pick_server(int(cfg.get("port", 8360)))
    port = server.server_address[1]
    url = f"http://127.0.0.1:{port}/"
    gui = args.gui or (_is_frozen and not args.no_window)

    if gui:
        print(f"[{APP_NAME}] 独立窗口模式 → {url}")
        serve_in_thread(server)
        try:
            run_gui(url)
        finally:
            server.shutdown()
        return

    print(f"[{APP_NAME}] 浏览器模式 → {url}  （Ctrl+C 停止）")
    if not args.no_browser:
        import webbrowser
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print(f"\n[{APP_NAME}] 已停止")


if __name__ == "__main__":
    main()
